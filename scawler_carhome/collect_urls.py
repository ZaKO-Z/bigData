"""
汽车之家车系URL自动采集器

从 https://www.autohome.com.cn/grade/carhtml/A.html ~ Z.html
提取所有车系ID，生成 config/series/{id}.html 格式的URL，
写入 series_urls.xlsx

用法: python collect_urls.py
"""
import sys
import re
import logging
from pathlib import Path

_BASE_DIR = Path(__file__).resolve().parent
if str(_BASE_DIR) not in sys.path:
    sys.path.insert(0, str(_BASE_DIR))

import requests
import openpyxl

from config import resolve_path, get_headers

logging.basicConfig(level=logging.INFO, format='%(asctime)s [%(levelname)s] %(message)s')
logger = logging.getLogger(__name__)

LETTER_PAGES = list('ABCDEFGHIJKLMNOPQRSTUVWXYZ')
BASE_URL = 'https://www.autohome.com.cn/grade/carhtml/{letter}.html'
OUTPUT_FILE = resolve_path('series_urls.xlsx')


def fetch_series_ids(letter):
    """从单个字母页面提取所有车系ID"""
    url = BASE_URL.format(letter=letter)
    try:
        resp = requests.get(url, headers=get_headers(), timeout=15)
        resp.raise_for_status()
        # 匹配 <li id="sXXXX"> 或 <li class="..." id="sXXXX">
        ids = set(re.findall(r'<li[^>]*\sid=\"s(\d+)\"', resp.text))
        logger.info("  %s: %d 个车系", letter, len(ids))
        return ids
    except requests.exceptions.RequestException as e:
        logger.warning("  %s: 请求失败 - %s", letter, e)
        return set()


def collect_all():
    """遍历A-Z，收集所有车系ID，去重后写入Excel"""
    all_ids = set()
    logger.info("开始采集车系URL（A-Z）...")

    for letter in LETTER_PAGES:
        ids = fetch_series_ids(letter)
        all_ids.update(ids)

    logger.info("共计去重后 %d 个唯一车系ID", len(all_ids))

    # 生成车系URL
    urls = sorted(
        f"https://www.autohome.com.cn/config/series/{sid}.html"
        for sid in all_ids
    )

    # 写入Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "车系URL"
    ws.cell(row=1, column=1, value="车系URL")
    for i, url in enumerate(urls, 2):
        ws.cell(row=i, column=1, value=url)

    # 列宽
    ws.column_dimensions['A'].width = 60

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)
    logger.info("已写入 %d 个URL → %s", len(urls), OUTPUT_FILE)

    # 预估产出
    avg_per_series = 8.5  # 平均每个车系约8.5个版本
    estimated = len(urls) * avg_per_series
    logger.info("预估产出: ~%.0f 条记录（按每车系 %.1f 个版本计）",
                 estimated, avg_per_series)

    return urls


if __name__ == '__main__':
    collect_all()
