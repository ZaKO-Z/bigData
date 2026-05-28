import sys
import asyncio
import logging
from pathlib import Path

# 确保包内模块可导入（兼容 python main.py 和 python -m 两种运行方式）
_BASE_DIR = Path(__file__).resolve().parent
if str(_BASE_DIR) not in sys.path:
    sys.path.insert(0, str(_BASE_DIR))

import openpyxl

from config import (
    INPUT_FILE, OUTPUT_CSV, CSV_FIELDS, FAILED_FILE, MAX_CONCURRENCY,
    COOKIE_COUNT, resolve_path,
)
from crawler import (
    CheckpointManager, FailedQueue, crawl_urls_async, crawl_urls,
)
from export import CsvAppender

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s [%(levelname)s] %(message)s',
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler(resolve_path('crawler.log'), encoding='utf-8'),
    ]
)
logger = logging.getLogger(__name__)


def read_urls_from_excel(filename=INPUT_FILE):
    file_path = resolve_path(filename)
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    urls = []
    for row in range(2, ws.max_row + 1):
        cell_value = ws.cell(row=row, column=1).value
        if cell_value and isinstance(cell_value, str) and cell_value.startswith('http'):
            urls.append(cell_value)
    return urls


def print_summary(total_versions, checkpoint, failed_queue):
    logger.info("=" * 50)
    logger.info("爬取完成!")
    logger.info("  本次新增版本: %d", total_versions)
    logger.info("  累计完成series: %d", len(checkpoint))
    if failed_queue:
        logger.warning("  失败URL: %d (详见 %s)", len(failed_queue), FAILED_FILE)
        for entry in failed_queue.failed:
            logger.warning("    - %s | %s", entry['series_id'], entry['reason'])
    logger.info("  CSV输出: %s", resolve_path(OUTPUT_CSV))
    logger.info("=" * 50)


async def amain(retry_failed=False, use_sync=False):
    # 断点恢复
    checkpoint = CheckpointManager()
    failed_queue = FailedQueue()

    # CSV增量写入器
    csv_appender = CsvAppender(OUTPUT_CSV, CSV_FIELDS)

    # 判断是全新开始还是续爬
    if checkpoint.completed:
        existing_rows = csv_appender.load_existing_count()
        logger.info("检测到断点: %d 个series已完成, CSV现有 %d 行数据",
                     len(checkpoint), existing_rows)
    else:
        csv_appender.init_header()
        logger.info("全新开始，已写入CSV表头")

    # 读取URL列表
    if retry_failed:
        urls = failed_queue.get_urls()
        if not urls:
            logger.info("失败队列为空，无需重试")
            return
        logger.info("重试失败队列: %d 个URL", len(urls))
        failed_queue.clear()
    else:
        urls = read_urls_from_excel()
        if not urls:
            logger.error("未找到有效的URL，请检查 %s", INPUT_FILE)
            return
        logger.info("共 %d 个URL", len(urls))

    # 爬取
    if use_sync:
        logger.info("使用同步模式（串行）")
        from crawler import CarDataCollector
        collector = CarDataCollector()
        crawl_urls(urls, collector)
        # 同步模式结束后写入CSV
        from export import to_csv
        to_csv(collector, OUTPUT_CSV, CSV_FIELDS)
        total_versions = len(collector)
    else:
        logger.info("使用异步模式（并发=%d, Cookie数=%d）", MAX_CONCURRENCY, COOKIE_COUNT)
        total_versions, _ = await crawl_urls_async(
            urls, checkpoint=checkpoint, failed_queue=failed_queue,
            csv_appender=csv_appender,
        )

    csv_appender.finish()
    print_summary(total_versions, checkpoint, failed_queue)


if __name__ == "__main__":
    retry_failed = '--retry-failed' in sys.argv
    use_sync = '--sync' in sys.argv
    asyncio.run(amain(retry_failed=retry_failed, use_sync=use_sync))
