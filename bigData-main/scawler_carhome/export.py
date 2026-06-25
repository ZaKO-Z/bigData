import csv
import logging
import sys
from pathlib import Path

_BASE_DIR = Path(__file__).resolve().parent
if str(_BASE_DIR) not in sys.path:
    sys.path.insert(0, str(_BASE_DIR))

import openpyxl
from openpyxl.styles import Font, Alignment

from config import resolve_path

logger = logging.getLogger(__name__)


class CsvAppender:
    """支持增量追加写入的CSV处理器"""

    def __init__(self, filename, fields):
        self.filename = filename
        self.fields = fields
        self._initialized = False
        self._row_count = 0

    @property
    def row_count(self):
        return self._row_count

    def init_header(self):
        """写入表头（仅首次调用生效）"""
        if self._initialized:
            return
        file_path = resolve_path(self.filename)
        file_path.parent.mkdir(parents=True, exist_ok=True)
        with open(file_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(self.fields)
        self._initialized = True
        logger.info("CSV表头已写入: %s", file_path)

    def append_rows(self, rows):
        """追加多行数据"""
        if not rows:
            return
        file_path = resolve_path(self.filename)
        with open(file_path, 'a', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            for row in rows:
                writer.writerow(row)
        self._row_count += len(rows)

    def load_existing_count(self):
        """读取已有CSV的数据行数（不含表头），用于checkpoint恢复时确认"""
        file_path = resolve_path(self.filename)
        if not file_path.exists():
            return 0
        with open(file_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            rows = list(reader)
        return max(0, len(rows) - 1)  # 减去表头

    def finish(self):
        logger.info("CSV写入完成，共 %d 行数据: %s",
                     self._row_count, resolve_path(self.filename))


def to_excel(collector, filename):
    """将收集器中的数据导出为Excel（全量导出，适用于中小规模）"""
    if not collector:
        raise ValueError("没有数据可保存")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汽车详细参数"

    basic_fields = ['品牌', '车系', '版本名称', '厂商指导价', 'seriesid', '年份款']
    param_fields = list(collector.params.keys())
    all_fields = basic_fields + param_fields

    for col, field in enumerate(all_fields, 1):
        cell = ws.cell(row=1, column=col, value=field)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')

    data_rows = len(collector)
    for row in range(2, data_rows + 2):
        idx = row - 2
        ws.cell(row=row, column=1, value=collector.brands[idx])
        ws.cell(row=row, column=2, value=collector.series[idx])
        ws.cell(row=row, column=3, value=collector.versions[idx])
        ws.cell(row=row, column=4, value=collector.prices[idx])
        ws.cell(row=row, column=5, value=collector.series_ids[idx])
        ws.cell(row=row, column=6, value=collector.years[idx])
        for col, param_name in enumerate(param_fields, 7):
            values = collector.params.get(param_name, [])
            value = values[idx] if idx < len(values) else ''
            ws.cell(row=row, column=col, value=value)

    for column in ws.columns:
        max_length = 0
        for cell in column:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column[0].column_letter].width = min(max_length + 2, 50)

    file_path = resolve_path(filename)
    file_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(file_path)
    logger.info("Excel已保存: %s (%d行)", file_path, data_rows)


def to_csv(collector, filename, fields):
    """全量导出到CSV"""
    if not collector:
        raise ValueError("没有数据可保存")

    file_path = resolve_path(filename)
    file_path.parent.mkdir(parents=True, exist_ok=True)

    with open(file_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(fields)
        for idx in range(len(collector)):
            row = [collector.get_param_value(field, idx) for field in fields]
            writer.writerow(row)

    logger.info("CSV已保存: %s (%d行)", file_path, len(collector))
